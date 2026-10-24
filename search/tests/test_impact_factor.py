"""
Tests for the journal impact factor: SIGAPS import, lookup, display and Excel export.
"""
import io
import json

import pytest
from django.core.management import call_command
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from search.services import impact_factor
from search.services.impact_factor import lookup_impact_factor
from search.services.participant_extractor import ParticipantExtractor
from search.services.outcome_extractor import OutcomeExtractor
from search.services.pubmed_client import _parse_article_xml
from search.tests.test_views import TEST_MIDDLEWARE


def _write_sigaps_workbook(path):
    """Minimal workbook with the layout of 'Rang SIGAPS 2020-2024.xlsx' (two header rows)."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'SIGAPS 2024'
    ws.append(['NLMid', 'ISSN', 'ESSN', 'Statut', 'Titre', 'JCR', 'Discipline(s)', 'Dernier IF',
               2020, None, 2021, None, 2022, None, 2023, None, 2024, None])
    ws.append([None, None, None, "d'indexation", None, 'Edition', 'WOS', 'connu',
               'IF', 'Cat.', 'IF', 'Cat.', 'IF', 'Cat.', 'IF', 'Cat.', 'IF', 'Cat.'])
    ws.append([372354, '0003-4932', '1528-1140', 'CI', 'Ann Surg', 'SCIE', 'YA', 2024,
               '9.0', 'A', '13.787', 'A', '10.1', 'A', '7.5', 'A', '6.5', 'A'])
    ws.append(['2985108R', '0017-5749', '1468-3288', 'CI', 'Gut', 'SCIE', 'KI', 2024,
               '23.1', 'A', '31.8', 'A', '24.5', 'A', '23', 'A', 26, 'A'])
    ws.append([101722685, None, '2474-9842', 'CI', 'BJS Open', 'SCIE', 'YA', 2024,
               None, None, None, None, '3.5', 'B', '3.5', 'B', '4,5', 'A'])
    ws.append([9716528, '0004-0010', None, 'CI', 'Arch Surg', 'SCIE', 'YA', 2012,
               None, None, None, None, None, None, None, None, None, None])  # no 2024 value
    wb.save(path)


@pytest.fixture
def if_file(tmp_path, settings):
    workbook = tmp_path / 'sigaps.xlsx'
    output = tmp_path / 'impact_factors_2024.json'
    _write_sigaps_workbook(workbook)
    settings.IMPACT_FACTOR_FILE = str(output)
    impact_factor.clear_cache()
    call_command('import_impact_factors', str(workbook), stdout=io.StringIO())
    yield output
    impact_factor.clear_cache()


class TestImport:
    def test_keeps_only_2024_values(self, if_file):
        data = json.loads(if_file.read_text(encoding='utf-8'))
        assert data['year'] == 2024
        by_ta = {j['ta']: j for j in data['journals']}
        assert by_ta['Ann Surg']['if'] == 6.5          # not the 2021 value (13.787)
        assert by_ta['Gut']['if'] == 26.0               # integer cell
        assert by_ta['BJS Open']['if'] == 4.5           # decimal comma
        assert 'Arch Surg' not in by_ta                 # no 2024 value: left out, no fallback year
        assert by_ta['Ann Surg']['issn'] == ['0003-4932', '1528-1140']

    def test_rejects_wrong_layout(self, tmp_path, settings):
        from django.core.management.base import CommandError
        wb = Workbook()
        wb.active.append(['foo', 'bar'])
        wb.active.append(['x', 'y'])
        bad = tmp_path / 'bad.xlsx'
        wb.save(bad)
        settings.IMPACT_FACTOR_FILE = str(tmp_path / 'out.json')
        with pytest.raises(CommandError):
            call_command('import_impact_factors', str(bad), stdout=io.StringIO())


class TestLookup:
    def test_nlm_id_with_pubmed_zero_padding(self, if_file):
        assert lookup_impact_factor(nlm_id='0372354') == {'impact_factor': 6.5, 'impact_factor_year': 2024}

    def test_nlm_id_with_letter(self, if_file):
        assert lookup_impact_factor(nlm_id='2985108R')['impact_factor'] == 26.0

    def test_issn_fallback(self, if_file):
        assert lookup_impact_factor(nlm_id='999', issns=['24749842'])['impact_factor'] == 4.5

    def test_abbreviation_fallback(self, if_file):
        assert lookup_impact_factor(journal_abbrev='ann surg')['impact_factor'] == 6.5

    def test_unknown_journal(self, if_file):
        assert lookup_impact_factor(nlm_id='123', issns=['0000-0000'], journal_abbrev='Nowhere J') is None

    def test_missing_file_returns_none(self, tmp_path, settings):
        settings.IMPACT_FACTOR_FILE = str(tmp_path / 'absent.json')
        impact_factor.clear_cache()
        assert lookup_impact_factor(nlm_id='0372354', journal_abbrev='Ann Surg') is None
        impact_factor.clear_cache()

    def test_file_created_after_first_lookup_is_used(self, tmp_path, settings):
        """Regression: a lookup made before the import must not keep the empty table forever."""
        path = tmp_path / 'late.json'
        settings.IMPACT_FACTOR_FILE = str(path)
        impact_factor.clear_cache()
        assert lookup_impact_factor(nlm_id='0372354') is None
        path.write_text(json.dumps({'year': 2024, 'journals': [
            {'nlm_id': '372354', 'issn': [], 'ta': 'Ann Surg', 'if': 6.5}]}), encoding='utf-8')
        assert lookup_impact_factor(nlm_id='0372354')['impact_factor'] == 6.5
        impact_factor.clear_cache()

    def test_ambiguous_issn_is_not_used(self, tmp_path, settings):
        path = tmp_path / 'ambiguous.json'
        path.write_text(json.dumps({'year': 2024, 'journals': [
            {'nlm_id': '1', 'issn': ['0218-5385'], 'ta': 'Asia Pac J Soc Work', 'if': 0.8},
            {'nlm_id': '2', 'issn': ['0218-5385'], 'ta': 'Asia Pac J Soc Work Dev', 'if': 1.2}]}), encoding='utf-8')
        settings.IMPACT_FACTOR_FILE = str(path)
        impact_factor.clear_cache()
        assert lookup_impact_factor(issns=['0218-5385']) is None
        assert lookup_impact_factor(issns=['0218-5385'], journal_abbrev='Asia Pac J Soc Work Dev')['impact_factor'] == 1.2
        impact_factor.clear_cache()


ABSTRACT = ("Background: We studied liver resection. Methods: A total of 240 patients were randomized. "
            "The primary outcome was 90-day mortality. Results: Mortality was lower.")

PUBMED_XML = f"""<?xml version="1.0"?>
<PubmedArticleSet><PubmedArticle><MedlineCitation>
  <PMID>12345</PMID>
  <Article>
    <Journal><ISSN IssnType="Electronic">1528-1140</ISSN><Title>Annals of surgery</Title></Journal>
    <ArticleTitle>A trial</ArticleTitle>
    <Abstract><AbstractText>{ABSTRACT}</AbstractText></Abstract>
    <AuthorList><Author><LastName>Doe</LastName><ForeName>J</ForeName>
      <AffiliationInfo><Affiliation>Department of Surgery, Lyon, France.</Affiliation></AffiliationInfo></Author></AuthorList>
    <Language>eng</Language>
  </Article>
  <MedlineJournalInfo><MedlineTA>Ann Surg</MedlineTA><NlmUniqueID>0372354</NlmUniqueID><ISSNLinking>0003-4932</ISSNLinking></MedlineJournalInfo>
</MedlineCitation></PubmedArticle></PubmedArticleSet>"""


class TestParsing:
    def test_journal_identifiers_and_impact_factor(self, if_file):
        article = _parse_article_xml(PUBMED_XML)[0]
        assert article['journal_abbrev'] == 'Ann Surg'
        assert article['nlm_unique_id'] == '0372354'
        assert article['issn'] == ['1528-1140', '0003-4932']
        assert article['impact_factor'] == 6.5
        assert article['impact_factor_year'] == 2024
        assert article['surgery_tier'] == 'high'

    def test_extracted_values_unchanged(self, if_file):
        """Adding journal fields must not change what the extractors return."""
        article = _parse_article_xml(PUBMED_XML)[0]
        assert article['sample_size'] == ParticipantExtractor.extract_sample_size(ABSTRACT)['sample_size']
        assert article['primary_outcome'] == OutcomeExtractor.extract_outcomes(ABSTRACT)['primary_outcome']

    def test_unknown_journal_has_no_impact_factor(self, if_file):
        xml = PUBMED_XML.replace('0372354', '0000001').replace('Ann Surg', 'Some J') \
                        .replace('1528-1140', '9999-9999').replace('0003-4932', '9999-9998')
        article = _parse_article_xml(xml)[0]
        assert article['impact_factor'] is None
        assert article['surgery_tier'] == ''


@pytest.mark.django_db
@override_settings(MIDDLEWARE=TEST_MIDDLEWARE)
class TestExcelExport(TestCase):
    def _export(self, articles):
        response = APIClient().post('/export/excel/', {'articles': articles}, format='json')
        assert response.status_code == 200
        ws = load_workbook(io.BytesIO(response.content)).active
        rows = list(ws.iter_rows(values_only=True))
        return rows[0], rows[1:]

    def test_columns_and_values_are_aligned(self):
        header, rows = self._export([{
            'pmid': '1', 'title': 'T', 'year': 2024, 'journal': 'Annals of surgery', 'language': 'fre',
            'impact_factor': 6.5, 'sample_size': 120,
        }])
        assert header[:7] == ('PMID', 'Title', 'Link', 'Publication Year', 'Journal', 'Impact Factor (2024)', 'Language')
        row = dict(zip(header, rows[0]))
        assert row['Journal'] == 'Annals of surgery'
        assert row['Language'] == 'French'
        assert row['Impact Factor (2024)'] == 6.5
        assert row['Sample Size'] == '120'

    def test_unknown_impact_factor_is_na(self):
        header, rows = self._export([{'pmid': '2', 'title': 'T', 'journal': 'Unknown J'}])
        assert dict(zip(header, rows[0]))['Impact Factor (2024)'] == 'NA'


@pytest.mark.django_db
def test_excel_export_looks_up_missing_impact_factor(if_file, client):
    response = client.post('/export/excel/', {'articles': [{'pmid': '3', 'title': 'T', 'journal': 'Gut',
                           'nlm_unique_id': '2985108R'}]}, content_type='application/json')
    ws = load_workbook(io.BytesIO(response.content)).active
    header, row = list(ws.iter_rows(values_only=True))[:2]
    assert dict(zip(header, row))['Impact Factor (2024)'] == 26.0


@pytest.mark.django_db
class TestPages:
    def test_search_page(self, if_file, client):
        html = client.get('/').content.decode()
        assert '>Journal</label>' in html and '>Rank</label>' not in html
        assert 'id="maxResults"' not in html
        assert 'Quality Assessment' not in html and 'Evidence Level' not in html and 'getJournalRanking' not in html
        assert html.count('class="surgery-tier-checkbox') == 2
        assert 'High Impact Factor' in html and 'Medium Impact Factor' in html
        assert 'IF 6.5' in html          # Annals of Surgery in the popup
        assert '23 high IF + 18 medium IF journals' in html
        assert '144 nursing journals' in html
        assert 'vendor/fontawesome/css/all.min.css' in html    # icons stylesheet, served locally
        for french in ('Non spécifié', 'Vider le cache', 'recuperation'):
            assert french not in html, french
        for removed in ('SIGAPS', 'zenodo', '/about/#validation'):
            assert removed not in html, removed
        assert 'Clarivate Journal Citation Reports' in html

    def test_about_page(self, client):
        html = client.get('/about/').content.decode()
        for expected in ('Clarivate Journal Citation Reports', 'not been measured',
                         'id="limits"', 'github.com/Daniblue25/Medical-Application'):
            assert expected in html, expected
        for removed in ('id="validation"', '#validation', 'id="cite"', 'How to cite', 'Zenodo', 'zenodo', 'SIGAPS', 'not validated',
                        'Experimental', 'unique in the world', 'The only tool in the world', '80%', '8 seconds',
                        '100 ms', 'RevMan', 'Rank A+', '76 unit tests', 'Confidence Score', 'Austere'):
            assert removed not in html, removed
